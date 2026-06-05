from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from backend.generators.base import BaseGenerator


class XlsxGenerator(BaseGenerator):
    def generate(self, prompt: str, template_path: Path | None = None) -> Path:
        data = self._llm_json(
            'Return JSON: {"title": "...", "headers": ["..."], "rows": [["..."]]}',
            prompt,
        )
        title = data.get("title", "Spreadsheet")
        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]
        headers = data.get("headers", [])
        if headers:
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="1B4F8A", end_color="1B4F8A", fill_type="solid")
        for row in data.get("rows", []):
            ws.append(row)
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)
        out = self._unique_output_path(title.replace(" ", "_"), "xlsx")
        wb.save(str(out))
        self._audit.log("file_generated", {"type": "xlsx", "path": str(out)}, category="DOCUMENT")
        return out
